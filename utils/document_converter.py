from docx import Document
from pdf2docx import Converter
from docx2pdf import convert
from fpdf import FPDF
import os

def convert_to_word(input_path, output_path):
    if input_path.endswith('.pdf'):
        # Convert PDF to Word
        cv = Converter(input_path)
        cv.convert(output_path, start=0, end=None)
        cv.close()
    elif input_path.endswith('.txt'):
        # Convert text to Word
        doc = Document()
        with open(input_path, 'r') as f:
            doc.add_paragraph(f.read())
        doc.save(output_path)
    else:
        raise ValueError("Unsupported input format for Word conversion")

def convert_to_pdf(input_path, output_path):
    if input_path.endswith('.docx'):
        # Convert Word to PDF
        convert(input_path, output_path)
    elif input_path.endswith('.txt'):
        # Convert text to PDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        
        with open(input_path, 'r') as f:
            for line in f:
                pdf.cell(200, 10, txt=line, ln=1)
        
        pdf.output(output_path)
    else:
        raise ValueError("Unsupported input format for PDF conversion")